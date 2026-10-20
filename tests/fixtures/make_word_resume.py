"""Build word_resume.docx, a fictional resume made the way Word users make them,
and export it to word_resume.pdf with LibreOffice.

What makes it Word-like: Times New Roman, Word's own bullet list style (the
bullet is a Symbol-font character), dates and places on a right-aligned tab
stop, headings with a bottom border, and real hyperlinks.

    uv run --with python-docx python tests/fixtures/make_word_resume.py

The person and every detail are made up.
"""

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).parent
WIDTH = Inches(7.0)  # the text width: letter paper with 0.75 in margins

doc = Document()
for section in doc.sections:
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.75)
    section.top_margin = section.bottom_margin = Inches(0.6)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)


def link(paragraph, text: str, url: str) -> None:
    """A hyperlink run, blue and underlined as Word draws it."""
    rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, value in (("w:color", "0563C1"), ("w:u", "single")):
        element = OxmlElement(tag)
        element.set(qn("w:val"), value)
        props.append(element)
    run.append(props)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def heading(text: str) -> None:
    """Bold capitals with a rule under them: a paragraph bottom border."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11.5)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def tabbed(left: list[tuple[str, str]], right: str, italic_right: bool = False) -> None:
    """A line with text on the left and `right` on a right-aligned tab stop.
    `left` is (text, style) pieces: style is "b", "i", "" or a URL."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    for text, style in left:
        if style.startswith("http"):
            link(paragraph, text, style)
        else:
            run = paragraph.add_run(text)
            run.bold, run.italic = style == "b", style == "i"
    if right:
        run = paragraph.add_run("\t" + right)
        run.italic = italic_right


def bullet(text: str) -> None:
    paragraph = doc.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(1)


name = doc.add_paragraph()
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = name.add_run("Priya Raman")
run.bold = True
run.font.size = Pt(20)
contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.add_run("Bengaluru, India | +91-90000-55555 | ")
link(contact, "priya.raman@example.com", "mailto:priya.raman@example.com")
contact.add_run(" | ")
link(contact, "linkedin.com/in/priya-raman-example", "https://linkedin.com/in/priya-raman-example")

heading("Professional Summary")
doc.add_paragraph(
    "Data analyst turned machine learning engineer with three years of shipping forecasting "
    "models for retail. Builds clean data pipelines and explains results to teams outside data. "
    "Looking for a role on a small applied ML team."
)

heading("Work Experience")
tabbed([("Larkspur Analytics", "b")], "Mar 2023 – Present")
tabbed([("Machine Learning Engineer", "i")], "Bengaluru, India", italic_right=True)
bullet(
    "Built a demand forecasting model in Python and LightGBM that cut stock-outs by 18% across "
    "40 stores, and set up weekly retraining with drift checks on every feature."
)
bullet("Moved nightly feature jobs from cron scripts to Airflow, with alerts on late or failed runs.")
bullet("Wrote the team's model review checklist, now used for every release.")
tabbed([("Northgate Retail", "b")], "Jul 2021 – Feb 2023")
tabbed([("Data Analyst", "i")], "Chennai, India", italic_right=True)
bullet("Automated the weekly sales report in SQL and Power BI, saving about six hours a week.")
bullet("Cleaned five years of point-of-sale data and documented every field for the analytics team.")

heading("Projects")
tabbed(
    [
        ("Shelf Monitor", "b"),
        (" | Python, OpenCV (", ""),
        ("live", "https://shelf-monitor.example.com"),
        (", ", ""),
        ("code", "https://github.com/priya-raman-example/shelf-monitor"),
        (")", ""),
    ],
    "2024",
)
bullet("Counted products on shelf photos with a fine-tuned YOLO model at 92% precision.")
tabbed(
    [("Route Planner", "b"), (" | Python, OR-Tools (", ""), ("live", "https://routes.example.org"), (")", "")],
    "2023",
)
bullet("Planned delivery routes for 30 vans, about 11% shorter than the manual plan.")

heading("Education")
tabbed([("Anna University", "b")], "2017 – 2021")
tabbed([("B.E., Computer Science | CGPA 8.4/10", "i")], "Chennai, India", italic_right=True)

heading("Skills")
for label, items in (("Languages:", " Python, SQL, R"), ("Tools:", " Airflow, Power BI, LightGBM, Git, Docker")):
    paragraph = doc.add_paragraph()
    paragraph.add_run(label).bold = True
    paragraph.add_run(items)

heading("Certifications")
doc.add_paragraph("Microsoft Certified: Azure Data Scientist Associate")

docx = HERE / "word_resume.docx"
doc.save(docx)
subprocess.run(
    ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(docx)],
    check=True,
    capture_output=True,
)
print("wrote", docx, "and", docx.with_suffix(".pdf"))
